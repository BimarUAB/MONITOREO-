import re
import sqlite3
import unicodedata
from pathlib import Path
from openpyxl import load_workbook

XLSX_DIR = Path(r'C:\Users\BIMAR\Desktop\UAB 2.2026\PRACTICA EMPRESARIAL\SECUNDARIA')
DB_PATH = r'C:\Users\BIMAR\Desktop\UAB 2.2026\PRACTICA EMPRESARIAL\LABORATORIOS\seguimiento-academico\data\app.db'


def curso_por_seccion(seccion: str):
    texto = str(seccion).strip().upper()
    match = re.match(r'^(\d+)([AB])$', texto)
    if not match:
        raise ValueError(f'Sección inválida: {seccion!r}')
    anio = match.group(1)
    letra = match.group(2)
    return f'{anio}° Secundaria {letra}'


def normalizar(texto):
    texto = unicodedata.normalize('NFD', str(texto or ''))
    texto = ''.join(c for c in texto if unicodedata.category(c) != 'Mn')
    return ' '.join(texto.upper().split())


def seccion_por_archivo(path):
    nombre = normalizar(path.stem)
    niveles = {
        'PRIMERO': '1', 'SEGUNDO': '2', 'TERCERO': '3',
        'CUARTO': '4', 'QUINTO': '5', 'SEXTO': '6',
    }
    nivel = next((numero for palabra, numero in niveles.items() if palabra in nombre), None)
    grupo = next((letra for letra in ('A', 'B') if re.search(rf'\b{letra}\b', nombre)), None)
    if not nivel or not grupo:
        raise ValueError(f'No se pudo identificar la sección en {path.name!r}')
    return f'{nivel}{grupo}'


con = sqlite3.connect(DB_PATH)
cur = con.cursor()

# Asegurar gestión activa
gestion = cur.execute('SELECT id, anio FROM gestiones WHERE activa = 1 LIMIT 1').fetchone()
if not gestion:
    raise RuntimeError('No hay una gestión activa en la base de datos.')

gestion_id, anio = gestion

# Asegurar cursos separados para cada sección A/B.
course_ids = {}
for anio_num in range(1, 7):
    for seccion in ('A', 'B'):
        nombre = f'{anio_num}° Secundaria {seccion}'
        cur.execute(
            'INSERT OR IGNORE INTO cursos (gestion_id, nombre, nivel, capacidad_max) VALUES (?, ?, ?, ?)',
            (gestion_id, nombre, 'Secundaria', 40),
        )
        course_ids[f'{anio_num}{seccion}'] = cur.execute(
            'SELECT id FROM cursos WHERE gestion_id = ? AND nombre = ?',
            (gestion_id, nombre),
        ).fetchone()[0]

creados = 0
actualizados = 0
saltados = 0
resumen = []

for workbook_path in sorted(XLSX_DIR.glob('*.xlsx')):
    current_course = seccion_por_archivo(workbook_path)
    curso_id = course_ids[current_course]
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    curso_creados = 0
    curso_actualizados = 0

    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[1] is None and row[2] is None and row[3] is None:
            continue

        apellido_p = str(row[1] or '').strip()
        apellido_m = str(row[2] or '').strip()
        nombres = str(row[3] or '').strip()
        apellidos = ' '.join(part for part in [apellido_p, apellido_m] if part).strip()
        clave = (normalizar(nombres), normalizar(apellidos))
        if not nombres:
            saltados += 1
            continue

        existentes = cur.execute(
            'SELECT id, nombres, apellidos FROM estudiantes WHERE curso_id = ?',
            (curso_id,),
        ).fetchall()
        coincidencia = next(
            (est for est in existentes if (normalizar(est[1]), normalizar(est[2])) == clave),
            None,
        )

        if coincidencia:
            if coincidencia[1] != nombres or coincidencia[2] != apellidos:
                cur.execute(
                    'UPDATE estudiantes SET nombres = ?, apellidos = ? WHERE id = ?',
                    (nombres, apellidos, coincidencia[0]),
                )
            curso_actualizados += 1
            continue

        # Reubicar registros creados por la importación anterior, que usaba
        # un solo curso por grado y no distinguía las secciones A/B.
        grado = current_course[0]
        cursos_anteriores = cur.execute(
            "SELECT c.id FROM cursos c WHERE c.gestion_id = ? AND c.nombre = ?",
            (gestion_id, f'{grado}° Secundaria'),
        ).fetchall()
        legado = None
        for (curso_anterior_id,) in cursos_anteriores:
            candidatos = cur.execute(
                'SELECT id, nombres, apellidos FROM estudiantes WHERE curso_id = ?',
                (curso_anterior_id,),
            ).fetchall()
            legado = next(
                (est for est in candidatos if (normalizar(est[1]), normalizar(est[2])) == clave),
                None,
            )
            if legado:
                break
        if legado:
            cur.execute(
                'UPDATE estudiantes SET nombres = ?, apellidos = ?, curso_id = ? WHERE id = ?',
                (nombres, apellidos, curso_id, legado[0]),
            )
            curso_actualizados += 1
            continue

        numero = int(row[0]) if isinstance(row[0], (int, float)) else curso_creados + 1
        rude = f'{anio}-{current_course}-{numero:03d}'
        while cur.execute('SELECT 1 FROM estudiantes WHERE rude = ?', (rude,)).fetchone():
            numero += 1
            rude = f'{anio}-{current_course}-{numero:03d}'
        cur.execute(
            'INSERT INTO estudiantes (rude, nombres, apellidos, curso_id) VALUES (?, ?, ?, ?)',
            (rude, nombres, apellidos, curso_id),
        )
        curso_creados += 1

    creados += curso_creados
    actualizados += curso_actualizados
    resumen.append((current_course, curso_creados, curso_actualizados))

con.commit()
print(f'IMPORTACIÓN COMPLETADA. Insertados: {creados}. Actualizados/verificados: {actualizados}. Saltados: {saltados}')
for curso, nuevos, verificados in sorted(resumen):
    print(f'{curso}: nuevos={nuevos}, verificados={verificados}')
con.close()
